from __future__ import annotations

import hashlib
import os
import subprocess
import tempfile
from pathlib import Path

SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".doc", ".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}
EQUIVALENT_EXTENSIONS = (".docx", ".txt", ".md", ".pdf")
EXTRACTOR_VERSION = "2"
PDF_MIN_TEXT_CHARS_BEFORE_OCR = 200
DEFAULT_OCR_MAX_PAGES = 0
PDF_OCR_HEADER_RATIO = 0.55


def _subprocess_creationflags() -> int:
    return getattr(subprocess, "CREATE_NO_WINDOW", 0)


def _env_int(name: str, default: int) -> int:
    try:
        return max(0, int(os.environ.get(name, default)))
    except ValueError:
        return default


def _ocr_enabled() -> bool:
    return os.environ.get("PHANLOAIVANBAN_OCR_ENABLED", "1").lower() not in {"0", "false", "no", "off"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_text(
    path: Path,
    runtime_dir: Path = Path("runtime"),
    *,
    full_text: bool = False,
    ocr_pages: int = 2,
) -> tuple[str, str, str | None]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore"), "ok", None
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pdf":
        return _extract_pdf(path, runtime_dir, full_text=full_text, ocr_pages=ocr_pages)
    if suffix == ".doc":
        return _extract_doc(path, runtime_dir)
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        if not _ocr_enabled():
            return "", "ocr_disabled", "tesseract"
        return _extract_image(path, runtime_dir, crop_header=True)
    return "", f"unsupported extension {suffix}", None


def _extract_docx(path: Path) -> tuple[str, str, str | None]:
    try:
        import docx  # type: ignore[import-not-found]
        from docx.table import Table  # type: ignore[import-not-found]
        from docx.text.paragraph import Paragraph  # type: ignore[import-not-found]
    except ImportError:
        return "", "python-docx not bundled", None
    document = docx.Document(str(path))
    lines: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, document).text
            if text.strip():
                lines.append(text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if paragraph.text.strip():
                            lines.append(paragraph.text)
    return "\n".join(lines), "ok", None

def _extract_pdf(
    path: Path,
    runtime_dir: Path = Path("runtime"),
    *,
    full_text: bool = False,
    ocr_pages: int = 2,
) -> tuple[str, str, str | None]:
    try:
        import pypdfium2 as pdfium  # type: ignore[import-not-found]
    except ImportError:
        return "", "missing_pdf_parser", None
    pdf = pdfium.PdfDocument(str(path))
    configured_pages = _env_int("PHANLOAIVANBAN_OCR_MAX_PAGES", ocr_pages)
    page_count = len(pdf) if full_text or configured_pages == 0 else min(len(pdf), max(1, configured_pages))
    pages = []
    for index in range(page_count):
        text_page = pdf[index].get_textpage()
        pages.append(text_page.get_text_range())
    text = "\n".join(pages).strip()
    min_text_chars = _env_int("PHANLOAIVANBAN_PDF_MIN_TEXT_CHARS_BEFORE_OCR", PDF_MIN_TEXT_CHARS_BEFORE_OCR)
    if len(text) >= min_text_chars:
        return text, "ok", None
    if not _ocr_enabled() or page_count == 0:
        return text, "ok" if text else "ocr_disabled", None
    ocr_text, ocr_status, ocr_engine = _extract_rendered_pdf_pages(pdf, page_count, runtime_dir)
    if len(ocr_text) > len(text):
        return ocr_text, ocr_status, ocr_engine
    if text:
        return text, "ok", None
    return ocr_text, ocr_status, ocr_engine


def _extract_rendered_pdf_pages(pdf, page_count: int, runtime_dir: Path) -> tuple[str, str, str | None]:
    last_status = "pdf_no_text"
    ocr_engine: str | None = None
    pages: list[str] = []
    with tempfile.TemporaryDirectory(prefix="phanloaivanban-pdf-ocr-") as temp_dir:
        output_dir = Path(temp_dir)
        for index in range(page_count):
            image_path = output_dir / f"page-{index + 1}.png"
            image = pdf[index].render(scale=2).to_pil()
            if hasattr(image, "crop") and hasattr(image, "size"):
                width, height = image.size
                image = image.crop((0, 0, width, int(height * PDF_OCR_HEADER_RATIO)))
            image.save(image_path)
            page_text, status, page_engine = _extract_image(image_path, runtime_dir)
            ocr_engine = page_engine or ocr_engine
            if status != "ok":
                last_status = status
                continue
            if page_text.strip():
                pages.append(page_text.strip())
    if pages:
        return "\n\n".join(pages), "ok", ocr_engine
    return "", last_status, ocr_engine

def _extract_image(path: Path, runtime_dir: Path, crop_header: bool = False) -> tuple[str, str, str | None]:
    tesseract = runtime_dir / "tesseract" / "tesseract.exe"
    if not tesseract.exists():
        return "", "missing_tesseract", "tesseract"
    traineddata = runtime_dir / "tesseract" / "tessdata" / "vie.traineddata"
    if not traineddata.exists():
        return "", "missing_tesseract_language", "tesseract"
    if crop_header:
        try:
            from PIL import Image  # type: ignore[import-not-found]
        except ImportError:
            return "", "missing_image_parser", "tesseract"
        with tempfile.TemporaryDirectory(prefix="phanloaivanban-image-ocr-") as temp_dir:
            try:
                image = Image.open(path)
            except Exception:
                return _extract_image(path, runtime_dir, crop_header=False)
            width, height = image.size
            cropped = image.crop((0, 0, width, int(height * PDF_OCR_HEADER_RATIO)))
            cropped_path = Path(temp_dir) / path.name
            cropped.save(cropped_path)
            return _extract_image(cropped_path, runtime_dir, crop_header=False)
    result = subprocess.run(
        [str(tesseract), str(path), "stdout", "-l", "vie", "--psm", "6"],
        creationflags=_subprocess_creationflags(),
        capture_output=True,
        encoding="utf-8",
        env={**os.environ, "TESSDATA_PREFIX": str(traineddata.parent)},
        errors="replace",
        text=True,
        timeout=120,
        check=False,
    )
    if result.returncode != 0:
        return "", "ocr_failed", "tesseract"
    return result.stdout, "ok", "tesseract"


def _extract_doc(path: Path, runtime_dir: Path) -> tuple[str, str, str | None]:
    soffice = runtime_dir / "libreoffice" / "program" / "soffice.exe"
    if not soffice.exists():
        equivalent = find_equivalent_document(path)
        if equivalent is not None:
            text, status, _ = extract_text(equivalent, runtime_dir)
            if text.strip() and status == "ok":
                return text, "ok", f"equivalent:{equivalent.name}"
        text = _extract_doc_utf16_text(path)
        if text.strip():
            return text, "ok", "doc-utf16"
        return "", "missing_libreoffice", None
    with tempfile.TemporaryDirectory(prefix="phanloaivanban-doc-") as temp_dir:
        output_dir = Path(temp_dir)
        result = subprocess.run(
            [str(soffice), "--headless", "--convert-to", "docx", "--outdir", str(output_dir), str(path)],
            creationflags=_subprocess_creationflags(),
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        converted = output_dir / f"{path.stem}.docx"
        if result.returncode != 0 or not converted.exists():
            equivalent = find_equivalent_document(path)
            if equivalent is not None:
                text, status, _ = extract_text(equivalent, runtime_dir)
                if text.strip() and status == "ok":
                    return text, "ok", f"equivalent:{equivalent.name}"
            return "", result.stderr.strip() or "DOC conversion failed", None
        return _extract_docx(converted)


def _extract_doc_utf16_text(path: Path) -> str:
    try:
        raw = path.read_bytes().decode("utf-16le", errors="ignore")
    except OSError:
        return ""
    import unicodedata

    text = raw.replace("\x07", "\n").replace("\r", "\n")
    allowed_punctuation = set(" \n\t.,;:()/_-–—“”\"'")
    text = "".join(
        char
        if char in allowed_punctuation or char.isdigit() or "LATIN" in unicodedata.name(char, "")
        else " "
        for char in text
    )
    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line and sum(ch.isalpha() or ch.isdigit() for ch in line) >= 2]
    for index, line in enumerate(lines[:200]):
        upper = line.upper()
        if "CỘNG HÒA" in upper or "QUỐC HỘI" in upper or "CHÍNH PHỦ" in upper or upper.startswith("SỐ"):
            return "\n".join(lines[index:])
    return "\n".join(lines)


def document_key_from_name(file_name: str) -> str | None:
    import re

    match = re.search(r"(\d{1,4})[_-](?:(\d{4})[_-])?(QH\d+|[A-ZĐ]{2,}(?:-[A-ZĐa-z]{2,})*)", file_name, flags=re.IGNORECASE)
    if not match:
        return None
    number, year, term = match.groups()
    year_part = f"_{year}" if year else ""
    return f"{int(number)}{year_part}_{term.upper()}"


def find_equivalent_document(path: Path) -> Path | None:
    candidates = find_equivalent_documents(path)
    return candidates[0] if candidates else None


def find_equivalent_documents(path: Path) -> list[Path]:
    key = document_key_from_name(path.name)
    if key is None:
        return []
    search_roots = [path.parent]
    parent = path.parent.parent
    if parent not in search_roots:
        search_roots.append(parent)
    if parent.exists():
        search_roots.extend(child for child in parent.iterdir() if child.is_dir() and child not in search_roots)
    candidates: list[Path] = []
    for root in search_roots:
        if not root.exists():
            continue
        for candidate in root.iterdir():
            if not candidate.is_file() or candidate == path:
                continue
            if candidate.suffix.lower() not in EQUIVALENT_EXTENSIONS:
                continue
            if document_key_from_name(candidate.name) == key:
                candidates.append(candidate)
    return sorted(candidates, key=lambda item: EQUIVALENT_EXTENSIONS.index(item.suffix.lower()))
